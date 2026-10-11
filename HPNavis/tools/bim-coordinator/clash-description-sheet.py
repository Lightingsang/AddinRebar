"""The two meeting sheets of the HP clash workbook, laid out like the company's BIM clash agenda:
  Summary            clashes per discipline x Mới / Chờ / Đã giải quyết, unresolved total, bar chart
  CLASH DESCRIPTION  one row per clash (New/Active/Reviewed/Resolved; Approved left out) with its Navisworks picture

Loaded by build-clash-report.py (importlib, the file name has hyphens). Pictures come from the Navisworks HTML report
of the same export (clash_pictures in clash-report-html.py); a clash without one says so in its photo cell.
"""

import io
import os

from PIL import Image as PilImage
from openpyxl.chart import BarChart, Reference, Series
from openpyxl.drawing.image import Image
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

NAVY = PatternFill("solid", fgColor="002060")
STATUS_FILL = PatternFill("solid", fgColor="F1C232")
WHITE_BOLD = Font(name="Arial", size=11, bold=True, color="FFFFFF")
BODY = Font(name="Arial", size=10)
THIN = Side(style="thin")
BOX = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
CENTER = Alignment(horizontal="center", vertical="center", wrap_text=True)
LEFT = Alignment(horizontal="left", vertical="center", wrap_text=True)

FIRST_ROW = 5
PICTURE_MAX_PX = (320, 180)
PICTURE_ROW_PT = 140
TEXT_ROW_PT = 66  # the description runs to four lines at Arial 10
DISCIPLINE_COLUMN = {"ARC": "M", "STR": "N", "MEP": "O"}
DISCIPLINE_NAME = {"ARC": "ARCHITECTURE", "STR": "STRUCTURE", "MEP": "MEPF"}
STATUS_COLUMN = {"New": "T", "Active": "T", "Reviewed": "V", "Resolved": "U"}
REPORTED = set(STATUS_COLUMN)
PRIORITY_LABEL = {1: "HIGH", 2: "MEDIUM", 3: "LOW"}
NO_PICTURE = "chưa có ảnh"

# (cell range, text) of the two header rows, as in the template
HEADERS = [
    ("A2:A3", "STT\nNo."), ("B2:B3", "Tình trạng\nStatus"), ("C2:C3", "Clash Photo Name"), ("D2:D3", "Vị trí\nLocation"),
    ("E2:E3", "Trục\nGrid"), ("F2:G2", "RFI"), ("F3", "No."), ("G3", "BIM"), ("H2:H3", "Mô tả xung đột\nClash Description"),
    ("I2:I3", "LINK Clash Photo"), ("J2:J3", "Bản vẽ được duyệt\nApproved drawing"), ("K2:K3", "Ngày phát hiện\nFound Date"),
    ("L2:L3", "Ý kiến của Chủ đầu tư\nClient's comment"), ("M2:O3", "Giải đáp bởi Tư vấn\nResponse by Consultants"),
    ("P2:P3", "Ngày phản hồi\nReplied Date"), ("Q2:Q3", "Bản vẽ thi công\nShop Drawing No."), ("R2:S2", "Mức ưu tiên\nPriority level"),
    ("R3", "Nhóm"), ("S3", "Mức"), ("T2:V2", "Tình trạng\nStatus"), ("T3", "Mới\nNew"), ("U3", "Xong\nDone"), ("V3", "Chờ\nPending"),
    ("W2:W3", "Thi công\nAs-Built"), ("M4", "Architecture"), ("N4", "Structure"), ("O4", "MEPF"),
]
WIDTHS = {"A": 7.7, "B": 8.3, "C": 48, "D": 17.3, "E": 13.3, "F": 6.1, "G": 6.7, "H": 49.4, "I": 22, "J": 20.7, "K": 16,
          "L": 36.4, "M": 13, "N": 13, "O": 13, "P": 16, "Q": 24.4, "R": 8, "S": 10, "T": 9.4, "U": 9.4, "V": 10, "W": 14}


def disciplines_of(pair):
    """ARC-STR -> ['ARC', 'STR']; a same-discipline pair counts once."""
    return list(dict.fromkeys(part for part in pair.split("-") if part in DISCIPLINE_COLUMN))


def reported_rows(snapshot):
    rows = [r for r in snapshot["rows"] if r["status"] in REPORTED]
    return sorted(rows, key=lambda r: (r["priority"], r["test"], natural(r["name"])))


def natural(name):
    digits = "".join(ch for ch in name if ch.isdigit())
    return (name.rstrip("0123456789"), int(digits) if digits else 0)


def describe(row):
    a, b = row["item1"], row["item2"]
    side = lambda item: f"{item['category'] or '?'} \"{item['name']}\"" + (f" (Id {item['elementId']})" if item["elementId"] else "")
    return (f"{row['name']} — {row['pair']} Nhóm {row['priority']}: {side(a)} ↔ {side(b)}; "
            f"khoảng cách {row['distanceMm']:g} mm\nTest: {row['test']}")


def header_cell(sheet, ref, text):
    first = ref.split(":")[0]
    if ":" in ref:
        sheet.merge_cells(ref)
    cell = sheet[first]
    cell.value = text
    for line in sheet[ref] if ":" in ref else [[cell]]:
        for c in line:
            c.fill, c.font, c.alignment, c.border = NAVY, WHITE_BOLD, CENTER, BOX


def picture_link(image_path, link_base):
    """Relative to the workbook's folder, so the link survives copying ClashReports; absolute when on another drive."""
    try:
        return os.path.relpath(image_path, link_base).replace(os.sep, "/")
    except ValueError:
        return image_path.resolve().as_uri()


def description_sheet(book, snapshot, pictures, link_base):
    """CLASH DESCRIPTION: returns (last data row, rows listed, rows without picture, pictured rows whose picture is
    missing). Found Date = the snapshot's foundAt (local), else the report's (UTC; older snapshots), else blank."""
    sheet = book.create_sheet("CLASH DESCRIPTION")
    rows = reported_rows(snapshot)
    last = max(FIRST_ROW, FIRST_ROW + len(rows) - 1)
    sheet.merge_cells("A1:Q1")
    sheet["A1"] = "CLASH DESCRIPTION"
    sheet["A1"].font, sheet["A1"].alignment = Font(name="Arial", size=18, bold=True), CENTER
    sheet["T1"] = "Clash Total"
    sheet["T1"].font = Font(name="Arial", size=11, bold=True)
    sheet["W1"] = f"=COUNT(A{FIRST_ROW}:A{last})"
    sheet["W1"].font = Font(name="Arial", size=18, bold=True, color="FF0000")
    for ref, text in HEADERS:
        header_cell(sheet, ref, text)
    for column in "ABCDEFGHIJKLPQRSTUVW":
        cell = sheet[f"{column}4"]
        cell.fill, cell.font, cell.border, cell.alignment = NAVY, WHITE_BOLD, BOX, CENTER
    sheet["A4"] = f"=COUNT(A{FIRST_ROW}:A{last})"
    sheet["B4"] = f"=SUM(B{FIRST_ROW}:B{last})"
    for column in "TUVW":
        sheet[f"{column}4"] = f"=SUM({column}{FIRST_ROW}:{column}{last})"
    sheet.row_dimensions[1].height = 43.5
    sheet.row_dimensions[3].height = 28.5
    for column, width in WIDTHS.items():
        sheet.column_dimensions[column].width = width
    sheet.freeze_panes = f"A{FIRST_ROW}"

    missing = lost = 0
    for index, row in enumerate(rows):
        r = FIRST_ROW + index
        picture = pictures.get((row["test"], row["name"]), {})
        values = {
            "A": index + 1, "D": row["level"], "E": picture.get("grid", ""), "H": describe(row),
            # the snapshot's local time first: the HTML report prints Date Found in UTC (live, THCSLT +7 h)
            "K": row.get("foundAt") or picture.get("found") or "", "R": row["priority"], "S": PRIORITY_LABEL.get(row["priority"], ""),
            STATUS_COLUMN[row["status"]]: 1,
        }
        for discipline in disciplines_of(row["pair"]):
            values[DISCIPLINE_COLUMN[discipline]] = 1
        for column, value in values.items():
            sheet[f"{column}{r}"] = value
        for column in WIDTHS:
            cell = sheet[f"{column}{r}"]
            cell.border, cell.font = BOX, BODY
            cell.alignment = LEFT if column in "CDEHL" else CENTER
        for column in "TUV":
            sheet[f"{column}{r}"].fill = STATUS_FILL
        image_path = picture.get("image")
        if image_path and image_path.exists():
            place_picture(sheet, image_path, f"C{r}")
            sheet[f"C{r}"] = image_path.name
            sheet[f"I{r}"] = image_path.name
            link = picture_link(image_path, link_base)
            sheet[f"I{r}"].hyperlink = link  # one string per cell: a shared Hyperlink object keeps only the last cell
            sheet[f"C{r}"].hyperlink = link
            sheet.row_dimensions[r].height = PICTURE_ROW_PT
        else:
            sheet[f"C{r}"] = NO_PICTURE
            sheet.row_dimensions[r].height = TEXT_ROW_PT
            missing += 1
            lost += 1 if row.get("pictured") else 0
    sheet.column_dimensions["I"].hidden = True
    return last, len(rows), missing, lost


def place_picture(sheet, path, anchor):
    """A thumbnail goes into the workbook (Navisworks renders 1024 px pictures, ~50 KB each); the cell links the original."""
    with PilImage.open(path) as original:
        thumbnail = original.convert("RGB")
        thumbnail.thumbnail(PICTURE_MAX_PX)
        data = io.BytesIO()
        thumbnail.save(data, format="JPEG", quality=85)
    data.seek(0)
    image = Image(data)
    sheet.add_image(image, anchor)


def summary_sheet(book, last, truncated=False):
    """Summary: Mới / Chờ / Đã giải quyết per discipline, counted from CLASH DESCRIPTION by formula."""
    sheet = book.create_sheet("Summary", 0)
    sheet.merge_cells("A1:K1")
    sheet["A1"] = "CLASH SUMMARY"
    sheet["A1"].font, sheet["A1"].alignment = Font(name="Arial", size=18, bold=True), CENTER
    headers = ["STT\nNo.", "Gói\nPackage", "Mới\nNew", "Đang chờ\nPending", "Đã giải quyết\nCompleted", "Tổng số\nTotal", "Đã ký\nSigned-off"]
    for column, text in zip("ABCDEFG", headers):
        header_cell(sheet, f"{column}2", text)
    header_cell(sheet, "H2:K2", "Số lượng xung đột chưa giải quyết\nTotal Unresolved Clashes")
    desc = "'CLASH DESCRIPTION'!"
    span = lambda column: f"{desc}${column}${FIRST_ROW}:${column}${last}"
    sheet["B3"] = "All"
    sheet["C3"], sheet["D3"], sheet["E3"] = f"={desc}T4", f"={desc}V4", f"={desc}U4"
    sheet["F3"], sheet["G3"] = "=SUM(C3:E3)", f"={desc}B4"
    sheet.merge_cells("H3:K6")
    sheet["H3"] = "=C3+D3"
    sheet["H3"].font, sheet["H3"].alignment = Font(name="Arial", size=28, bold=True, color="FF0000"), CENTER
    for offset, discipline in enumerate(DISCIPLINE_COLUMN):
        r = 4 + offset
        own = DISCIPLINE_COLUMN[discipline]
        sheet[f"A{r}"], sheet[f"B{r}"] = offset + 1, DISCIPLINE_NAME[discipline]
        sheet[f"C{r}"] = f"=COUNTIFS({span(own)},1,{span('T')},1)"
        sheet[f"D{r}"] = f"=COUNTIFS({span(own)},1,{span('V')},1)"
        sheet[f"E{r}"] = f"=COUNTIFS({span(own)},1,{span('U')},1)"
        sheet[f"F{r}"] = f"=SUM(C{r}:E{r})"
        sheet[f"G{r}"] = f"=COUNTIFS({span(own)},1,{span('B')},1)"
    for line in sheet["A3:G6"]:
        for cell in line:
            cell.border, cell.font, cell.alignment = BOX, Font(name="Arial", size=11, bold=cell.row == 3), CENTER
    sheet["A8"] = "Mỗi clash được đếm cho cả hai bộ môn của nó; dòng All đếm mỗi clash một lần."
    sheet["A8"].font = Font(name="Arial", size=9, italic=True)
    if truncated:
        sheet["A9"] = "Snapshot bị cắt ở 50 000 dòng: Summary và CLASH DESCRIPTION đếm thiếu; số đúng theo test ở sheet Tổng hợp."
        sheet["A9"].font = Font(name="Arial", size=9, bold=True, color="FF0000")
    for column, width in zip("ABCDEFGHIJK", [7, 18, 11, 11, 13, 10, 10, 8, 8, 8, 8]):
        sheet.column_dimensions[column].width = width
    sheet.row_dimensions[2].height = 45

    chart = BarChart()
    chart.type, chart.grouping, chart.overlap = "col", "clustered", 0
    chart.title, chart.y_axis.title = "Clash theo bộ môn", "Số clash"
    for column, title in zip((3, 4, 5), ("Mới", "Đang chờ", "Đã giải quyết")):  # disciplines only, not the All row
        chart.series.append(Series(Reference(sheet, min_col=column, min_row=4, max_row=6), title=title))
    chart.set_categories(Reference(sheet, min_col=2, min_row=4, max_row=6))
    chart.x_axis.delete = chart.y_axis.delete = False  # recent Excel hides openpyxl axes unless told otherwise
    chart.width, chart.height = 18, 9
    sheet.add_chart(chart, "A11")
    return sheet
