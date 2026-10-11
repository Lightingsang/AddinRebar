"""Turn an HP clash snapshot (written by the seed bim_export_clash_snapshot) into the coordination workbook.

The snapshot is ClashReports/clash-LOD<phase>-<yyyyMMdd-HHmmss>.json beside the model. The workbook gets six sheets:
  Summary            per discipline: Mới / Chờ / Đã giải quyết (the company's BIM clash agenda layout, with a chart)
  CLASH DESCRIPTION  one row per New/Active/Reviewed/Resolved clash with its Navisworks picture (the agenda layout)
  Tổng hợp   results by priority group (Nhóm) x discipline pair x status, with the change since the previous export
  Theo tầng  results by level x Nhóm x status
  Clash      one row per clash (test, status, distance, centre, both elements, level)
  Thay đổi   clashes new since the previous export and clashes resolved or gone since then
The previous export is the newest older snapshot of the same phase in the same folder (or --previous). A result is
matched by its GUID; when the GUID is new, by rule + both Element Ids (Navisworks keeps a result's GUID while the test is
re-run, so the fallback only catches tests that were deleted and recreated).

Pictures come from the Navisworks HTML report written with the snapshot (its "htmlReport" folder) or from --html,
e.g. a report the coordinator wrote by hand in Clash Detective > Report (HTML (Tabular), all tests combined).

Usage:
    python build-clash-report.py SNAPSHOT.json [--previous OLDER.json] [--html REPORT.html] [--out REPORT.xlsx]
    python build-clash-report.py --latest FOLDER          (newest snapshot in FOLDER)
"""

import argparse
import importlib.util
import json
import sys
import time
from collections import Counter, defaultdict
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

STATUSES = ["New", "Active", "Reviewed", "Approved", "Resolved"]
OPEN_STATUSES = {"New", "Active", "Reviewed"}
GROUP_NAMES = {1: "Nhóm 1 - HIGH", 2: "Nhóm 2 - MEDIUM", 3: "Nhóm 3 - LOW"}
HEADER_FILL = PatternFill("solid", fgColor="0696D7")
HEADER_FONT = Font(bold=True, color="FFFFFF")
SAVE_ATTEMPTS = 4
HERE = Path(__file__).resolve().parent


def _sibling(file_name):
    """Load a helper module of this folder (their file names have hyphens, so no plain import)."""
    spec = importlib.util.spec_from_file_location(file_name.replace("-", "_")[:-3], HERE / file_name)
    module = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(module)
    return module


description = _sibling("clash-description-sheet.py")
report_html = _sibling("clash-report-html.py")


def fail(message):
    print(f"ERROR: {message}", file=sys.stderr)
    sys.exit(1)


def load(path):
    data = json.loads(Path(path).read_text(encoding="utf-8"))
    if data.get("schemaVersion") != 1 or "rows" not in data or "tests" not in data:
        fail(f"{path} is not an HP clash snapshot (schemaVersion 1)")
    return data


def snapshots_in(folder, phase=None):
    files = sorted(Path(folder).glob("clash-LOD*.json"))
    return [f for f in files if phase is None or f.name.startswith(f"clash-LOD{phase}-")]


def find_previous(snapshot_path, phase):
    older = [f for f in snapshots_in(snapshot_path.parent, phase) if f.name < snapshot_path.name]
    return older[-1] if older else None


def key_of(row):
    return (row["ruleId"], row["item1"].get("elementId", ""), row["item2"].get("elementId", ""))


def compare(current, previous):
    """New = in current, unseen before. Resolved = open before and now Resolved or no longer reported."""
    if previous is None:
        return None
    before_by_guid = {r["guid"]: r for r in previous["rows"]}
    before_by_key = {key_of(r): r for r in previous["rows"]}
    seen = set()
    new, resolved = [], []
    for row in current["rows"]:
        before = before_by_guid.get(row["guid"]) or before_by_key.get(key_of(row))
        if before is None:
            new.append(row)
            continue
        seen.add(before["guid"])
        if before["status"] in OPEN_STATUSES and row["status"] == "Resolved":
            resolved.append((row, "Resolved"))
    current_tests = {t["name"] for t in current["tests"]}
    for row in previous["rows"]:
        if row["guid"] not in seen and row["status"] in OPEN_STATUSES and row["test"] in current_tests:
            resolved.append((row, "không còn trong kết quả"))
    return {"new": new, "resolved": resolved}


def style_header(sheet, row=1):
    for cell in sheet[row]:
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(vertical="center", wrap_text=True)
    sheet.freeze_panes = f"A{row + 1}"  # a string, so no empty cell is created before the first append


def fit_columns(sheet, widths):
    for index, width in enumerate(widths, start=1):
        sheet.column_dimensions[get_column_letter(index)].width = width


def summary_sheet(book, current, previous, changes):
    sheet = book.active
    sheet.title = "Tổng hợp"
    tests = current["tests"]
    not_run = sum(1 for t in tests if not t.get("lastRun"))
    sheet.append([f"Mô hình: {current['model']}  |  Giai đoạn LOD{current['phase']} @ {current['toleranceMm']:g} mm  |  xuất lúc {current['exportedAt']}"])
    sheet.append([f"{len(tests)} test, {not_run} chưa chạy  |  so với: {previous['exportedAt'] if previous else 'lần đầu (chưa có lần xuất trước)'}"])
    sheet.append([])
    header = ["Nhóm", "Cặp bộ môn", "Số test"] + STATUSES + ["Tổng", "Tổng lần trước", "Chênh lệch", "Mới", "Đã xử lý"]
    sheet.append(header)
    style_header(sheet, 4)

    def counts(snapshot):
        table = defaultdict(Counter)
        for test in snapshot["tests"]:
            for status, n in test["byStatus"].items():
                table[(test["priority"], test["pair"])][status] += n
        return table

    now, before = counts(current), counts(previous) if previous else None
    tests_per = Counter((t["priority"], t["pair"]) for t in tests)
    new_per = Counter((r["priority"], r["pair"]) for r in changes["new"]) if changes else Counter()
    resolved_per = Counter((r["priority"], r["pair"]) for r, _ in changes["resolved"]) if changes else Counter()
    for key in sorted(tests_per):
        priority, pair = key
        total = sum(now[key].values())
        line = [GROUP_NAMES.get(priority, f"P{priority}"), pair, tests_per[key]] + [now[key].get(s, 0) for s in STATUSES] + [total]
        if before is None:
            line += ["", "", "", ""]
        else:
            previous_total = sum(before[key].values())
            line += [previous_total, total - previous_total, new_per[key], resolved_per[key]]
        sheet.append(line)
    total_line = ["Tổng", "", len(tests)] + [sum(now[k].get(s, 0) for k in now) for s in STATUSES] + [sum(sum(c.values()) for c in now.values())]
    if before is not None:
        previous_total = sum(sum(c.values()) for c in before.values())
        total_line += [previous_total, total_line[-1] - previous_total, len(changes["new"]), len(changes["resolved"])]
    sheet.append(total_line)
    for cell in sheet[sheet.max_row]:
        cell.font = Font(bold=True)
    fit_columns(sheet, [18, 12, 9] + [10] * len(STATUSES) + [9, 14, 11, 8, 10])


def level_sheet(book, current):
    sheet = book.create_sheet("Theo tầng")
    sheet.append(["Tầng", "Nhóm"] + STATUSES + ["Tổng"])
    style_header(sheet)
    table = defaultdict(Counter)
    for row in current["rows"]:
        table[(row["level"], row["priority"])][row["status"]] += 1
    for (level, priority) in sorted(table, key=lambda k: (k[0], k[1])):
        counts = table[(level, priority)]
        sheet.append([level, GROUP_NAMES.get(priority, f"P{priority}")] + [counts.get(s, 0) for s in STATUSES] + [sum(counts.values())])
    sheet.auto_filter.ref = sheet.dimensions
    fit_columns(sheet, [24, 18] + [10] * len(STATUSES) + [8])


ROW_HEADER = ["Test", "Nhóm", "Cặp bộ môn", "Clash", "Nhóm kết quả", "Trạng thái", "Khoảng cách (mm)", "X (mm)", "Y (mm)", "Z (mm)", "Tầng",
              "Cấu kiện 1", "Category 1", "Element Id 1", "File 1", "Cấu kiện 2", "Category 2", "Element Id 2", "File 2", "GUID"]


def row_values(row):
    a, b = row["item1"], row["item2"]
    x, y, z = (row["centerMm"] + [None, None, None])[:3]
    return [row["test"], GROUP_NAMES.get(row["priority"], f"P{row['priority']}"), row["pair"], row["name"], row.get("group") or "", row["status"],
            row["distanceMm"], x, y, z, row["level"], a["name"], a["category"], a["elementId"], a["sourceFile"],
            b["name"], b["category"], b["elementId"], b["sourceFile"], row["guid"]]


def clash_sheet(book, current):
    sheet = book.create_sheet("Clash")
    sheet.append(ROW_HEADER)
    style_header(sheet)
    for row in current["rows"]:
        sheet.append(row_values(row))
    sheet.auto_filter.ref = sheet.dimensions
    fit_columns(sheet, [44, 16, 10, 9, 14, 11, 10, 12, 12, 10, 18, 26, 18, 12, 30, 26, 18, 12, 30, 38])


def changes_sheet(book, changes):
    sheet = book.create_sheet("Thay đổi")
    if changes is None:
        sheet.append(["Lần đầu: chưa có lần xuất trước cùng giai đoạn để so sánh."])
        return
    sheet.append(["Thay đổi"] + ROW_HEADER)
    style_header(sheet)
    for row in changes["new"]:
        sheet.append(["Mới"] + row_values(row))
    for row, how in changes["resolved"]:
        sheet.append([f"Đã xử lý ({how})"] + row_values(row))
    sheet.auto_filter.ref = sheet.dimensions
    fit_columns(sheet, [22, 44, 16, 10, 9, 14, 11, 10, 12, 12, 10, 18, 26, 18, 12, 30, 26, 18, 12, 30, 38])


def save(book, out):
    for attempt in range(1, SAVE_ATTEMPTS + 1):
        try:
            book.save(out)
            return
        except PermissionError:
            if attempt == SAVE_ATTEMPTS:
                fail(f"{out} is locked (open in Excel or held by Google Drive sync) — close it and run again")
            time.sleep(attempt)


def html_for(snapshot_path, snapshot, html=None):
    """The HTML report to take pictures from: the one given, else the one the export wrote, else none."""
    if html:
        return Path(html)
    if snapshot.get("htmlReport"):
        return snapshot_path.parent / snapshot["htmlReport"].replace("\\", "/")
    return None


def build(snapshot_path, previous_path=None, out=None, html=None):
    snapshot_path = Path(snapshot_path)
    current = load(snapshot_path)
    html_path = html_for(snapshot_path, current, html)
    pictures, picture_warnings = report_html.clash_pictures(html_path) if html_path else ({}, [])
    previous_path = Path(previous_path) if previous_path else find_previous(snapshot_path, current["phase"])
    previous = load(previous_path) if previous_path else None
    if previous and previous["phase"] != current["phase"]:
        fail(f"{previous_path.name} is phase {previous['phase']}, {snapshot_path.name} is {current['phase']}")
    changes = compare(current, previous)
    book = Workbook()
    summary_sheet(book, current, previous, changes)
    level_sheet(book, current)
    clash_sheet(book, current)
    changes_sheet(book, changes)
    out = Path(out) if out else snapshot_path.with_suffix(".xlsx")
    last, listed, missing, lost = description.description_sheet(book, current, pictures, out.parent)
    book.move_sheet("CLASH DESCRIPTION", offset=-(len(book.sheetnames) - 1))
    description.summary_sheet(book, last, current.get("truncated", False))
    book.active = 0
    book.calculation.fullCalcOnLoad = True  # Summary and the counters are formulas; Excel computes them on open
    save(book, out)
    return {"out": out, "previous": previous_path, "rows": len(current["rows"]), "listed": listed, "withoutPicture": missing,
            "html": html_path, "pictureWarnings": picture_warnings + ([f"{lost} clash(es) pictured by Navisworks were not found in the HTML report"] if lost else []),
            "new": len(changes["new"]) if changes else None, "resolved": len(changes["resolved"]) if changes else None}


def main():
    parser = argparse.ArgumentParser(description=__doc__.splitlines()[0])
    parser.add_argument("snapshot", nargs="?")
    parser.add_argument("--latest", help="folder: use its newest clash snapshot")
    parser.add_argument("--previous")
    parser.add_argument("--html", help="Navisworks HTML (Tabular) report to take the pictures from")
    parser.add_argument("--out")
    options = parser.parse_args()
    if options.latest:
        found = snapshots_in(options.latest)
        if not found:
            fail(f"no clash-LOD*.json in {options.latest}")
        options.snapshot = max(found, key=lambda f: f.stat().st_mtime)
    if not options.snapshot:
        parser.error("give a snapshot file or --latest FOLDER")
    result = build(options.snapshot, options.previous, options.out, options.html)
    compared = f"so với {Path(result['previous']).name}: {result['new']} mới, {result['resolved']} đã xử lý" if result["previous"] else "lần đầu"
    print(f"{result['rows']} clash -> {result['out']} ({compared})")
    print(f"CLASH DESCRIPTION: {result['listed']} dòng, {result['listed'] - result['withoutPicture']} có ảnh"
          + (f" (ảnh từ {result['html']})" if result["html"] else " (không có report HTML)"))
    for warning in result["pictureWarnings"]:
        print(f"WARNING: {warning}", file=sys.stderr)


if __name__ == "__main__":
    main()
