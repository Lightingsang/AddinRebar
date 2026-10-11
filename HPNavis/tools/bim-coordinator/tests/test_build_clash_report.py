"""Tests of build-clash-report.py on hand-made snapshots (no Navisworks).

Run from the repository root:
    python -m unittest discover -s HPNavis/tools/bim-coordinator/tests
"""

import importlib.util
import json
import tempfile
import unittest
from pathlib import Path

from openpyxl import load_workbook
from PIL import Image as PilImage

HERE = Path(__file__).resolve().parent
SPEC = importlib.util.spec_from_file_location("build_clash_report", HERE.parent / "build-clash-report.py")
report = importlib.util.module_from_spec(SPEC)
SPEC.loader.exec_module(report)

TEST_A = "HP|LOD350-400|P1|MEP-STR|M4 Pipes and Pipe Accessories-S3 Structural Columns"
TEST_B = "HP|LOD350-400|P2|ARC-STR|A5 Doors-S2 Structural Walls"


def item(name, element_id, level=""):
    return {"name": name, "category": "Pipes", "elementId": element_id, "sourceFile": "x.rvt", "level": level}


def row(test, rule, priority, pair, guid, status, level="TẦNG 2", ids=("1", "2")):
    return {"test": test, "ruleId": rule, "priority": priority, "pair": pair, "guid": guid, "name": "Clash1", "group": None,
            "status": status, "distanceMm": -12.5, "centerMm": [1, 2, 3], "level": level,
            "item1": item("Pipe", ids[0], level), "item2": item("Column", ids[1])}


def snapshot(rows, exported, tests=None):
    tests = tests or [
        {"name": TEST_A, "ruleId": "HP_M4_S3", "priority": 1, "pair": "MEP-STR", "left": "M4", "right": "S3", "status": "Complete",
         "lastRun": exported, "resultCount": 0, "byStatus": {}},
        {"name": TEST_B, "ruleId": "HP_A5_S2", "priority": 2, "pair": "ARC-STR", "left": "A5", "right": "S2", "status": "New",
         "lastRun": None, "resultCount": 0, "byStatus": {}},
    ]
    for test in tests:
        mine = [r for r in rows if r["test"] == test["name"]]
        test["resultCount"] = len(mine)
        test["byStatus"] = {}
        for r in mine:
            test["byStatus"][r["status"]] = test["byStatus"].get(r["status"], 0) + 1
    return {"schemaVersion": 1, "exportedAt": exported, "model": "THCSLT.nwd", "phase": "350-400", "toleranceMm": 10,
            "tests": tests, "rows": rows, "truncated": False}


class BuildClashReportTests(unittest.TestCase):
    def setUp(self):
        self.folder = Path(tempfile.mkdtemp())

    def write(self, name, data):
        path = self.folder / name
        path.write_text(json.dumps(data, ensure_ascii=False), encoding="utf-8")
        return path

    def test_first_export_has_every_sheet_and_says_it_is_the_first(self):
        path = self.write("clash-LOD350-400-20261011-090000.json", snapshot([
            row(TEST_A, "HP_M4_S3", 1, "MEP-STR", "g1", "New"),
            row(TEST_A, "HP_M4_S3", 1, "MEP-STR", "g2", "New", level=""),
        ], "2026-10-11T02:00:00Z"))
        result = report.build(path)
        book = load_workbook(result["out"])
        self.assertEqual(["Summary", "CLASH DESCRIPTION", "Tổng hợp", "Theo tầng", "Clash", "Thay đổi"], book.sheetnames)
        self.assertEqual(2, book["Clash"].max_row - 1)
        self.assertIn("Lần đầu", book["Thay đổi"]["A1"].value)
        self.assertIn("1 chưa chạy", book["Tổng hợp"]["A2"].value)
        summary_total = [r for r in book["Tổng hợp"].iter_rows(values_only=True) if r and r[0] == "Tổng"][0]
        self.assertEqual(2, summary_total[3 + len(report.STATUSES)])
        self.assertIsNone(result["previous"])

    def test_second_export_finds_new_and_resolved_clashes_by_guid(self):
        self.write("clash-LOD350-400-20261011-090000.json", snapshot([
            row(TEST_A, "HP_M4_S3", 1, "MEP-STR", "g1", "New"),
            row(TEST_A, "HP_M4_S3", 1, "MEP-STR", "g2", "Active", ids=("3", "4")),
            row(TEST_A, "HP_M4_S3", 1, "MEP-STR", "g3", "New", ids=("5", "6")),
        ], "2026-10-11T02:00:00Z"))
        later = self.write("clash-LOD350-400-20261018-090000.json", snapshot([
            row(TEST_A, "HP_M4_S3", 1, "MEP-STR", "g1", "Active"),
            row(TEST_A, "HP_M4_S3", 1, "MEP-STR", "g2", "Resolved", ids=("3", "4")),
            row(TEST_A, "HP_M4_S3", 1, "MEP-STR", "g9", "New", ids=("7", "8")),
        ], "2026-10-18T02:00:00Z"))
        result = report.build(later)
        self.assertEqual("clash-LOD350-400-20261011-090000.json", Path(result["previous"]).name)
        self.assertEqual(1, result["new"])        # g9
        self.assertEqual(2, result["resolved"])   # g2 resolved, g3 gone
        kinds = sorted(r[0] for r in load_workbook(result["out"])["Thay đổi"].iter_rows(min_row=2, values_only=True))
        self.assertEqual(["Mới", "Đã xử lý (Resolved)", "Đã xử lý (không còn trong kết quả)"], kinds)

    def test_a_recreated_test_is_matched_by_rule_and_element_ids(self):
        self.write("clash-LOD350-400-20261011-090000.json", snapshot([row(TEST_A, "HP_M4_S3", 1, "MEP-STR", "old-guid", "New")], "2026-10-11T02:00:00Z"))
        later = self.write("clash-LOD350-400-20261018-090000.json", snapshot([row(TEST_A, "HP_M4_S3", 1, "MEP-STR", "new-guid", "New")], "2026-10-18T02:00:00Z"))
        result = report.build(later)
        self.assertEqual(0, result["new"])
        self.assertEqual(0, result["resolved"])

    def test_an_unchanged_second_export_reports_nothing_new(self):
        rows = [row(TEST_A, "HP_M4_S3", 1, "MEP-STR", "g1", "Active")]
        self.write("clash-LOD350-400-20261011-090000.json", snapshot(rows, "2026-10-11T02:00:00Z"))
        later = self.write("clash-LOD350-400-20261018-090000.json", snapshot([dict(r) for r in rows], "2026-10-18T02:00:00Z"))
        result = report.build(later)
        self.assertEqual((0, 0), (result["new"], result["resolved"]))

    def test_a_snapshot_of_another_phase_is_never_the_previous_one(self):
        other = snapshot([], "2026-10-10T02:00:00Z")
        other["phase"] = "200"
        self.write("clash-LOD200-20261010-090000.json", other)
        later = self.write("clash-LOD350-400-20261018-090000.json", snapshot([], "2026-10-18T02:00:00Z"))
        self.assertIsNone(report.build(later)["previous"])

    def test_levels_group_the_clashes(self):
        path = self.write("clash-LOD350-400-20261011-090000.json", snapshot([
            row(TEST_A, "HP_M4_S3", 1, "MEP-STR", "g1", "New", level="TẦNG 2"),
            row(TEST_A, "HP_M4_S3", 1, "MEP-STR", "g2", "New", level="TẦNG 2"),
            row(TEST_B, "HP_A5_S2", 2, "ARC-STR", "g3", "New", level="TẦNG 3"),
        ], "2026-10-11T02:00:00Z"))
        rows = list(load_workbook(report.build(path)["out"])["Theo tầng"].iter_rows(min_row=2, values_only=True))
        self.assertEqual([("TẦNG 2", "Nhóm 1 - HIGH", 2), ("TẦNG 3", "Nhóm 2 - MEDIUM", 1)], [(r[0], r[1], r[-1]) for r in rows])


HTML_REPORT = r"""  # raw: Navisworks writes Windows paths<html><body>
<table class="testSummaryTable"><tr class="headerRow"><td rowspan="2" colspan="2" class="testName">{test}</td>
<td class="headerCell">Tolerance</td></tr><tr class="contentRow"><td class="contentCell">10.000mm</td></tr></table><br>
<table class="mainTable">
<tr class="headerRow"><td colspan="12" class="generalHeader"></td><td colspan="3" class="item1HeaderGroup">Item 1</td></tr>
<tr class="headerRow"><td colspan="2" class="generalHeader">Image</td><td colspan="2" class="generalHeader">Clash Name</td>
<td class="generalHeader">Status</td><td class="generalHeader">Distance</td><td class="generalHeader">Grid Location</td>
<td class="generalHeader">Priority</td><td class="generalHeader">Date Found</td><td colspan="3" class="generalHeader">Clash Point</td>
<td class="item1Header">Item ID</td></tr>
<tr class="contentRow"><td colspan="2" class="contentCell"><a target="_blank" href="clash-report_files\cd010001.jpg ">
<img border="0" width="95" height="95" src="clash-report_files\cd010001.jpg"></a></td><td colspan="2" class="contentCell">Clash1</td>
<td class="contentCell">New</td><td class="contentCell">-91.716</td><td class="contentCell">A-4 : TẦNG 2</td>
<td class="contentCell">0</td><td class="contentCell">2026/10/11 09:55</td><td colspan="3" class="contentCell">x:1</td>
<td class="item1Content"><i>Element ID</i>: 1</td></tr>
</table></body></html>"""


class PictureTests(unittest.TestCase):
    def setUp(self):
        self.folder = Path(tempfile.mkdtemp())

    def export(self, rows):
        stem = "clash-LOD350-400-20261011-090000"
        html = self.folder / stem / "clash-report.html"
        (html.parent / "clash-report_files").mkdir(parents=True)
        html.write_text(HTML_REPORT.format(test=TEST_A), encoding="utf-8")
        PilImage.new("RGB", (1024, 1024), "red").save(html.parent / "clash-report_files" / "cd010001.jpg")
        data = snapshot(rows, "2026-10-11T02:00:00Z")
        data["htmlReport"] = f"{stem}\\clash-report.html"  # the engine writes a Windows path
        path = self.folder / f"{stem}.json"
        path.write_text(json.dumps(data, ensure_ascii=False), encoding="utf-8")
        return path

    def test_the_navisworks_picture_lands_in_its_own_row(self):
        path = self.export([row(TEST_A, "HP_M4_S3", 1, "MEP-STR", "g1", "New")])
        result = report.build(path)
        sheet = load_workbook(result["out"])["CLASH DESCRIPTION"]
        self.assertEqual((1, 0), (result["listed"], result["withoutPicture"]))
        self.assertEqual(1, len(sheet._images))
        self.assertEqual(4, sheet._images[0].anchor._from.row)  # zero-based: row 5, the first clash
        self.assertEqual("cd010001.jpg", sheet["C5"].value)
        self.assertEqual("A-4 : TẦNG 2", sheet["E5"].value)
        self.assertEqual("2026/10/11 09:55", sheet["K5"].value)
        self.assertEqual("clash-LOD350-400-20261011-090000/clash-report_files/cd010001.jpg", sheet["I5"].hyperlink.target)

    def test_a_clash_without_picture_says_so_and_both_disciplines_count_it(self):
        path = self.export([row(TEST_A, "HP_M4_S3", 1, "MEP-STR", "g1", "New"), row(TEST_A, "HP_M4_S3", 1, "MEP-STR", "g2", "Active")])
        rows = report.load(path)["rows"]
        rows[1]["name"] = "Clash2"
        path.write_text(json.dumps({**report.load(path), "rows": rows}, ensure_ascii=False), encoding="utf-8")
        sheet = load_workbook(report.build(path)["out"])["CLASH DESCRIPTION"]
        self.assertEqual("chưa có ảnh", sheet["C6"].value)
        self.assertEqual((None, 1, 1), (sheet["M6"].value, sheet["N6"].value, sheet["O6"].value))  # MEP-STR: Structure + MEPF

    def test_statuses_map_to_the_agenda_columns_and_approved_is_left_out(self):
        rows = [row(TEST_A, "HP_M4_S3", 1, "MEP-STR", f"g{i}", status, ids=(str(i), "x")) for i, status in
                enumerate(["New", "Active", "Reviewed", "Resolved", "Approved"])]
        for index, r in enumerate(rows):
            r["name"] = f"Clash{index + 1}"
        sheet = load_workbook(report.build(self.export(rows))["out"])["CLASH DESCRIPTION"]
        flags = [tuple(sheet[f"{c}{r}"].value for c in "TUV") for r in range(5, 9)]
        self.assertEqual([(1, None, None), (1, None, None), (None, None, 1), (None, 1, None)], flags)
        self.assertIsNone(sheet["A9"].value)

    def test_a_main_table_without_its_own_summary_never_borrows_another_tests_picture(self):
        path = self.export([row(TEST_A, "HP_M4_S3", 1, "MEP-STR", "g1", "New")])
        html = path.with_suffix("") / "clash-report.html"
        page = html.read_text(encoding="utf-8")
        start = page.index('<table class="mainTable">')
        html.write_text(page.replace("</body>", page[start:page.index("</table></body>") + 8].replace("Clash1", "Clash9") + "</body>"), encoding="utf-8")
        pictures, warnings = report.report_html.clash_pictures(html)
        self.assertEqual([(TEST_A, "Clash1")], list(pictures))
        self.assertEqual([], warnings)

    def test_a_picture_navisworks_wrote_but_the_html_lost_is_warned_and_the_found_date_falls_back(self):
        rows = [row(TEST_A, "HP_M4_S3", 1, "MEP-STR", "g1", "New"), row(TEST_A, "HP_M4_S3", 1, "MEP-STR", "g2", "New", ids=("3", "4"))]
        rows[1]["name"], rows[1]["pictured"], rows[1]["foundAt"] = "Clash2", True, "2026/10/10 08:00"
        result = report.build(self.export(rows))
        self.assertEqual(1, result["withoutPicture"])
        self.assertTrue(any("not found in the HTML report" in w for w in result["pictureWarnings"]))
        self.assertEqual("2026/10/10 08:00", load_workbook(result["out"])["CLASH DESCRIPTION"]["K6"].value)


if __name__ == "__main__":
    unittest.main()
