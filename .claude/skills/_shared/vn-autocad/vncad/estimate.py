# -*- coding: utf-8 -*-
"""Lap bang du toan chi phi xay dung ra file Excel chuan doanh nghiep VN.

Quy tac bat buoc: 100% phong chu Times New Roman, cot so can phai dinh dang #,##0,
cong thuc dong (=E*F, =SUM), chieu cao dong tu co gian, vien bang mau xam nha nhan.
"""

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

FONT = "Times New Roman"
THIN = Side(style="thin", color="A6A6A6")
BOX = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
HEAD_FILL = PatternFill("solid", fgColor="D9D9D9")
SEC_FILL = PatternFill("solid", fgColor="F2F2F2")

COLS = [("STT", 6), ("Ma hieu", 11), ("Noi dung cong tac", 52), ("DVT", 8),
        ("Khoi luong", 12), ("Don gia (VND)", 15), ("Thanh tien (VND)", 18)]


def build_estimate(path, project, rows, vat=0.08, overhead=0.06, contingency=0.05):
    """Xuat file .xlsx du toan.

    project: dict {ten, dia_diem, chu_dau_tu, ngay}
    rows: danh sach, moi phan tu la
          ("SECTION", "PHAN I: PHAN NGAM VA MONG")  hoac
          (ma_hieu, noi_dung, dvt, khoi_luong, don_gia)
    """
    wb = Workbook()
    ws = wb.active
    ws.title = "DuToan_ChiPhi"
    _title(ws, project)
    last = _body(ws, rows)
    _summary(ws, last, vat, overhead, contingency)
    _layout(ws)
    wb.save(path)
    return path


# ---------------------------------------------------------------------------
def _title(ws, p):
    ws.merge_cells("A1:G1")
    ws["A1"] = "BANG TONG HOP DU TOAN CHI PHI XAY DUNG CONG TRINH"
    ws["A1"].font = Font(name=FONT, size=14, bold=True)
    ws["A1"].alignment = Alignment(horizontal="center", vertical="center")

    ws.merge_cells("A2:G2")
    ws["A2"] = "DU AN: %s  |  DIA DIEM: %s" % (p.get("ten", ""), p.get("dia_diem", ""))
    ws["A2"].font = Font(name=FONT, size=11, italic=True)
    ws["A2"].alignment = Alignment(horizontal="center")

    ws.merge_cells("A3:G3")
    ws["A3"] = "Chu dau tu: %s     Ngay lap: %s" % (p.get("chu_dau_tu", ""), p.get("ngay", ""))
    ws["A3"].font = Font(name=FONT, size=11, italic=True)
    ws["A3"].alignment = Alignment(horizontal="center")

    for i, (name, _) in enumerate(COLS, 1):
        c = ws.cell(row=5, column=i, value=name)
        c.font = Font(name=FONT, size=11, bold=True)
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        c.fill = HEAD_FILL
        c.border = BOX


def _body(ws, rows):
    r = 6
    stt = 0
    for item in rows:
        if item[0] == "SECTION":
            ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=7)
            c = ws.cell(row=r, column=1, value=item[1])
            c.font = Font(name=FONT, size=11, bold=True)
            c.fill = SEC_FILL
            c.alignment = Alignment(horizontal="left", vertical="center")
            for i in range(1, 8):
                ws.cell(row=r, column=i).border = BOX
            r += 1
            continue
        stt += 1
        ma, noi_dung, dvt, kl, dg = item
        _cell(ws, r, 1, stt, "center")
        _cell(ws, r, 2, ma, "center")
        _cell(ws, r, 3, noi_dung, "left", wrap=True)
        _cell(ws, r, 4, dvt, "center")
        _cell(ws, r, 5, kl, "right", fmt="#,##0.00")
        _cell(ws, r, 6, dg, "right", fmt="#,##0")
        _cell(ws, r, 7, "=E%d*F%d" % (r, r), "right", fmt="#,##0")
        r += 1
    return r


def _summary(ws, r, vat, overhead, contingency):
    first, last = 6, r - 1
    lines = [
        ("CONG CHI PHI TRUC TIEP", "=SUM(G%d:G%d)" % (first, last), True),
        ("Chi phi chung va quan ly (%d%%)" % round(overhead * 100), "=G%d*%s" % (r, overhead), False),
        ("Chi phi du phong (%d%%)" % round(contingency * 100), "=G%d*%s" % (r, contingency), False),
        ("Cong truoc thue", "=G%d+G%d+G%d" % (r, r + 1, r + 2), False),
        ("Thue GTGT (%d%%)" % round(vat * 100), "=G%d*%s" % (r + 3, vat), False),
        ("TONG CONG GIA TRI DU TOAN SAU THUE", "=G%d+G%d" % (r + 3, r + 4), True),
    ]
    for i, (label, formula, bold) in enumerate(lines):
        row = r + i
        ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=6)
        c = ws.cell(row=row, column=1, value=label)
        c.font = Font(name=FONT, size=11, bold=bold)
        c.alignment = Alignment(horizontal="right", vertical="center")
        v = ws.cell(row=row, column=7, value=formula)
        v.font = Font(name=FONT, size=11, bold=bold)
        v.number_format = "#,##0"
        v.alignment = Alignment(horizontal="right")
        for col in range(1, 8):
            ws.cell(row=row, column=col).border = BOX


def _cell(ws, r, c, value, align, fmt=None, wrap=False):
    cell = ws.cell(row=r, column=c, value=value)
    cell.font = Font(name=FONT, size=11)
    cell.alignment = Alignment(horizontal=align, vertical="center", wrap_text=wrap)
    cell.border = BOX
    if fmt:
        cell.number_format = fmt
    return cell


def _layout(ws):
    for i, (_, w) in enumerate(COLS, 1):
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.row_dimensions[1].height = 24
    ws.row_dimensions[5].height = 30
    ws.freeze_panes = "A6"
    ws.page_setup.orientation = "landscape"
    ws.page_setup.fitToWidth = 1
