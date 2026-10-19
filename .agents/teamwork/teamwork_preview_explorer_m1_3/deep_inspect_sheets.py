import olefile
import openpyxl
import sys
import io

# Ensure UTF-8 output
out_file = r'g:\09-PROJECT AI\01_Revit\02_CshapRevit\01_AddinRebar\.agents\teamwork\teamwork_preview_explorer_m1_3\excel_sheets_dump.txt'
with open(out_file, 'w', encoding='utf-8') as out:
    def log(msg=""):
        out.write(msg + "\n")

    log("=== INSPECTING Iso_QHPlus.xls WITH OLEFILE ===")
    ole = olefile.OleFileIO(r'C:\kata_pro\Iso_QHPlus.xls')
    log(f"OLE streams: {ole.listdir()}")
    if ole.exists('Workbook'):
        stream = ole.openstream('Workbook').read()
        log(f"Found 'Workbook' stream. Size: {len(stream):,} bytes")
        sheet_names = []
        pos = 0
        while pos < len(stream) - 4:
            rec_id = int.from_bytes(stream[pos:pos+2], 'little')
            rec_len = int.from_bytes(stream[pos+2:pos+4], 'little')
            if rec_id == 0x0085: # BOUNDSHEET
                cch = stream[pos+10]
                flag = stream[pos+11]
                if flag == 0 and pos + 12 + cch <= len(stream):
                    name = stream[pos+12:pos+12+cch].decode('latin-1', errors='ignore')
                    sheet_names.append(name)
                elif flag == 1 and pos + 12 + cch*2 <= len(stream):
                    name = stream[pos+12:pos+12+cch*2].decode('utf-16le', errors='ignore')
                    sheet_names.append(name)
            pos += 4 + rec_len
        log(f"Iso_QHPlus.xls Sheets: {sheet_names}")
    
    # Check VBA streams in Iso_QHPlus.xls
    vba_streams = [s for s in ole.listdir() if '_VBA_PROJECT_CUR' in s]
    log(f"Iso_QHPlus.xls VBA Streams count: {len(vba_streams)}")

    log("\n=== DEEP INSPECTING Kata.xlsm (Dam, Cot, Vach, Thong_ke) ===")
    wb = openpyxl.load_workbook(r'C:\kata_pro\Kata.xlsm', data_only=False)

    def inspect_sheet(ws, max_r=50, max_c=25):
        log(f"\n========================================================")
        log(f"SHEET: {ws.title} (Max Row: {ws.max_row}, Max Col: {ws.max_column})")
        log(f"========================================================")
        for r in range(1, max_r + 1):
            row_vals = []
            for c in range(1, max_c + 1):
                val = ws.cell(row=r, column=c).value
                if val is not None:
                    coord = ws.cell(row=r, column=c).coordinate
                    # truncate very long formula if needed
                    val_str = str(val)
                    if len(val_str) > 50:
                        val_str = val_str[:47] + "..."
                    row_vals.append(f"{coord}='{val_str}'")
            if row_vals:
                log(f"Row {r:2d}: " + " | ".join(row_vals))

    for sheet_name in ['Dam', 'Cot', 'Vach', 'Thong_ke', 'ThepDamAllMBKC', 'To_hop_thep']:
        if sheet_name in wb.sheetnames:
            inspect_sheet(wb[sheet_name], max_r=60, max_c=30)
        else:
            log(f"Sheet {sheet_name} not found!")

    wb.close()
    log("\n=== COMPLETED EXCEL INSPECTION ===")

print("Done writing to excel_sheets_dump.txt")
