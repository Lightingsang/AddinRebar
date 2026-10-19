import openpyxl, sys

wb = openpyxl.load_workbook('C:/kata_pro/Kata.xlsm', data_only=True)
with open('.agents/teamwork/teamwork_preview_spec_miner_survey_2/kata_xlsm_data.txt', 'w', encoding='utf-8') as out:
    for name in wb.sheetnames:
        sheet = wb[name]
        out.write(f"==================================================\n")
        out.write(f"SHEET: {name} (dims: {sheet.dimensions})\n")
        out.write(f"==================================================\n")
        for r in range(1, min(sheet.max_row + 1, 100)):
            row_vals = [sheet.cell(r, c).value for c in range(1, min(sheet.max_column + 1, 30))]
            if any(v is not None for v in row_vals):
                out.write(f"Row {r:3d}: {row_vals}\n")
        out.write("\n\n")

print("Dumped Kata.xlsm sheets.")
