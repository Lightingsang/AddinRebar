import openpyxl
import sys

sys.stdout.reconfigure(encoding='utf-8')

wb = openpyxl.load_workbook(r'C:\kata_pro\Kata.xlsm', data_only=False)
ws = wb['In_tem_thep']

print("=== TEMPLATE A1:F12 ===")
for r in range(1, 13):
    cells = [f"{openpyxl.utils.get_column_letter(c)}{r}='{ws.cell(r, c).value}'" for c in range(1, 7) if ws.cell(r, c).value is not None]
    print(f"Row {r:2d}: " + " | ".join(cells))

print("\n=== SAMPLE GENERATED TAG K1:P12 ===")
for r in range(1, 13):
    cells = [f"{openpyxl.utils.get_column_letter(c)}{r}='{ws.cell(r, c).value}'" for c in range(11, 17) if ws.cell(r, c).value is not None]
    print(f"Row {r:2d}: " + " | ".join(cells))
