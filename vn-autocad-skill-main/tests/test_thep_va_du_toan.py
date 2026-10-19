# -*- coding: utf-8 -*-
"""Kiem thu bang thong ke thep va viec du toan lay so tu ban ve.

So tong da soat tay ngay 10/09/2026. Doi thiet ke thi cap nhat cac hang so
duoi day; test la chot chan de khong ai vo tinh lam lech bang thong ke.
"""

from openpyxl import load_workbook

from examples.nha_pho_5x18m import build_dutoan, rebar_quantities as rq
from examples.nha_pho_5x18m.draw_stair_structure import REBAR_ROWS_THANG
from examples.nha_pho_5x18m.draw_structure import LOAI_SAN, REBAR_ROWS
from examples.nha_pho_5x18m.foundation_rebar import REBAR_ROWS_MONG
from vncad.struct import bar_weight

KG_TANG_DIEN_HINH = 2282.1
KG_DOAN_THANG = 123.7
KG_MONG = 748.9
KG_KHUNG_TOAN_NHA = 7217.6


def _kg(rows):
    return sum(bar_weight(d, dai / 1000.0 * sl * so) for _, _, _, d, dai, sl, so in rows)


def test_cong_thuc_trong_luong_thep():
    # W = D^2 / 162.2 x L; thep D16 nang 1,578 kg moi met
    assert abs(bar_weight(16, 1.0) - 1.578) < 0.001
    assert abs(bar_weight(10, 1.0) - 0.617) < 0.001


def test_thep_san_du_hai_phuong_va_thep_mu():
    # loi cu: bang chi ghi mot phuong, thieu ca thep mu
    san = [r for r in REBAR_ROWS if r[0].startswith("San")]
    assert len(san) == 3 * len(LOAI_SAN)
    hinh = {r[2] for r in san}
    assert {"Lop duoi phuong X", "Lop duoi phuong Y", "Mu goi bon canh"} <= hinh


def test_hieu_thanh_danh_so_lien_tuc():
    for rows in (REBAR_ROWS, REBAR_ROWS_THANG, REBAR_ROWS_MONG):
        assert [r[1] for r in rows] == [str(i) for i in range(1, len(rows) + 1)]


def test_tong_thep_khop_so_da_soat():
    assert abs(_kg(REBAR_ROWS) - KG_TANG_DIEN_HINH) < 0.2
    assert abs(_kg(REBAR_ROWS_THANG) - KG_DOAN_THANG) < 0.2
    assert abs(_kg(REBAR_ROWS_MONG) - KG_MONG) < 0.2
    nho, lon = rq.thep_than_nha()
    assert abs(nho + lon - KG_KHUNG_TOAN_NHA) < 0.2


def test_ham_luong_thep_khung_hop_ly():
    # nha pho khung nhe nhip 5m: khoang 20 den 45 kg thep moi m2 san
    nho, lon = rq.thep_than_nha()
    assert 20.0 <= (nho + lon) / 270.0 <= 45.0


def test_du_toan_lay_dung_so_tu_ban_ve(tmp_path):
    path = build_dutoan.build(str(tmp_path))
    ws = load_workbook(path).active
    dong = {r[1]: r for r in ws.iter_rows(min_row=6, values_only=True) if r[1]}
    nho, lon = rq.thep_than_nha()
    assert dong["AF.612"][4] == round(rq.thep_mong(), 1)
    assert dong["AF.622"][4] == round(lon, 1)
    assert dong["AF.623"][4] == round(nho, 1)
    for ma in ("AF.612", "AF.622", "AF.623"):
        assert "bang thong ke" in dong[ma][2]          # ghi ro nguon ngay trong dong


def test_du_toan_toan_phong_times(tmp_path):
    ws = load_workbook(build_dutoan.build(str(tmp_path))).active
    phong = {c.font.name for row in ws.iter_rows() for c in row if c.value is not None}
    assert phong == {"Times New Roman"}
